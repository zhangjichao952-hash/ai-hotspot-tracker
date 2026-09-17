# -*- coding: utf-8 -*-
"""
AI热点信息与短视频选题报告生成器
用法:
  python3 generate_report.py --output /path/to/output.xlsx --data /path/to/data.json

data.json 格式:
{
  "hot_topics": [
    [序号, 热点事件, 核心内容摘要, 发布时间, 讨论度, 普通用户可理解性, 画面可演示性, 同质化程度, 普通用户理解门槛, 综合评分, 主要信息来源],
    ...
  ],
  "topics": [
    [选题编号, 选题名称, 内容切口, 目标受众, 预计时长, 画面可演示性, 信息来源, 待核查项],
    ...
  ],
  "sources": [
    [来源网站, 网站定位, 覆盖的热点编号, 原始链接/搜索索引],
    ...
  ]
}

也可不传 --data，直接修改本脚本底部的 hot_topics / topics / sources 变量。
"""
import argparse, json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ============ 通用样式 ============
HEADER_FONT = Font(name='Arial', bold=True, size=11, color='FFFFFF')
HEADER_FILL = PatternFill('solid', start_color='2F5496')
CELL_FONT = Font(name='Arial', size=10)
WRAP_ALIGN = Alignment(wrap_text=True, vertical='top')
CENTER_ALIGN = Alignment(wrap_text=True, vertical='center', horizontal='center')
THIN_BORDER = Border(
    left=Side(style='thin', color='B4C6E7'),
    right=Side(style='thin', color='B4C6E7'),
    top=Side(style='thin', color='B4C6E7'),
    bottom=Side(style='thin', color='B4C6E7'),
)

def style_header_row(ws, row, ncols):
    for col in range(1, ncols + 1):
        c = ws.cell(row=row, column=col)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = CENTER_ALIGN
        c.border = THIN_BORDER

def style_data_cell(ws, row, col, value):
    c = ws.cell(row=row, column=col, value=value)
    c.font = CELL_FONT
    c.alignment = WRAP_ALIGN
    c.border = THIN_BORDER
    return c


def build_sheet1(wb, hot_topics):
    """Sheet1: 热点信息去重与评估"""
    ws = wb.active
    ws.title = '热点信息去重与评估'

    headers = [
        '序号', '热点事件', '核心内容摘要', '发布时间', '讨论度',
        '普通用户可理解性', '画面可演示性', '同质化程度', '普通用户理解门槛',
        '综合评分(1-5)', '主要信息来源'
    ]
    for i, h in enumerate(headers, 1):
        style_data_cell(ws, 1, i, h)
    style_header_row(ws, 1, len(headers))

    for row_idx, row_data in enumerate(hot_topics, 2):
        for col_idx, val in enumerate(row_data, 1):
            style_data_cell(ws, row_idx, col_idx, val)

    col_widths = [6, 22, 50, 14, 10, 22, 22, 28, 28, 12, 30]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'


def build_sheet2(wb, topics):
    """Sheet2: 短视频选题方案"""
    ws = wb.create_sheet('短视频选题方案')

    headers = ['选题编号', '选题名称', '内容切口（角度）', '目标受众', '预计时长', '画面可演示性', '信息来源', '待核查项']
    for i, h in enumerate(headers, 1):
        style_data_cell(ws, 1, i, h)
    style_header_row(ws, 1, len(headers))

    for row_idx, row_data in enumerate(topics, 2):
        for col_idx, val in enumerate(row_data, 1):
            style_data_cell(ws, row_idx, col_idx, val)

    col_widths = [8, 25, 55, 22, 12, 22, 35, 40]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'


def build_sheet3(wb, sources):
    """Sheet3: 信息来源对照"""
    ws = wb.create_sheet('信息来源对照')

    headers = ['来源网站', '网站定位', '覆盖的热点编号', '原始链接/搜索索引']
    for i, h in enumerate(headers, 1):
        style_data_cell(ws, 1, i, h)
    style_header_row(ws, 1, len(headers))

    for row_idx, row_data in enumerate(sources, 2):
        for col_idx, val in enumerate(row_data, 1):
            style_data_cell(ws, row_idx, col_idx, val)

    col_widths = [22, 35, 25, 35]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'


def main():
    parser = argparse.ArgumentParser(description='AI热点信息与短视频选题报告生成器')
    parser.add_argument('--output', '-o', default='AI行业热点信息与短视频选题.xlsx', help='输出Excel文件路径')
    parser.add_argument('--data', '-d', default=None, help='JSON数据文件路径（可选）')
    args = parser.parse_args()

    if args.data and os.path.exists(args.data):
        with open(args.data, 'r', encoding='utf-8') as f:
            data = json.load(f)
        hot_topics = data.get('hot_topics', [])
        topics = data.get('topics', [])
        sources = data.get('sources', [])
    else:
        # 使用脚本内嵌的示例数据（替换为实际数据）
        hot_topics = HOT_TOPICS
        topics = TOPICS
        sources = SOURCES

    wb = Workbook()
    build_sheet1(wb, hot_topics)
    build_sheet2(wb, topics)
    build_sheet3(wb, sources)
    wb.save(args.output)
    print(f'文件已保存到: {args.output}')


# ============ 示例数据（替换为实际搜集的数据）============
# 每行格式: [序号, 热点事件, 核心内容摘要, 发布时间, 讨论度, 普通用户可理解性, 画面可演示性, 同质化程度, 普通用户理解门槛, 综合评分, 主要信息来源]
HOT_TOPICS = [
    [1, '示例热点事件', '这里是核心内容摘要...', '2026-09-01', '高', '中', '高', '中', '低门槛', 4, '示例来源'],
]

# 每行格式: [选题编号, 选题名称, 内容切口, 目标受众, 预计时长, 画面可演示性, 信息来源, 待核查项]
TOPICS = [
    [1, '示例选题', '从什么角度切入...', '科技爱好者', '3-5分钟', '高', 'web_xxxx', '需要核查的具体问题...'],
]

# 每行格式: [来源网站, 网站定位, 覆盖的热点编号, 原始链接/搜索索引]
SOURCES = [
    ['示例来源', '示例定位', '1', 'example.com'],
]


if __name__ == '__main__':
    main()
